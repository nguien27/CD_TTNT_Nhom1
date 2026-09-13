# -*- coding: utf-8 -*-
"""GĐ1-02: Thử nghiệm đọc DOCX gồm paragraph và table."""

from __future__ import annotations

import json
from pathlib import Path

THU_MUC_DU_AN = Path(__file__).resolve().parents[2]
TEST_FILE = THU_MUC_DU_AN / "test_data" / "gd1_02" / "test_nhan_vien.docx"


def doc_docx(duong_dan: Path):
    from docx import Document

    doc = Document(duong_dan)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    tables = []
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        rows = []
        for row in table.rows[1:]:
            values = [cell.text.strip() for cell in row.cells]
            if not any(values):
                continue
            rows.append(dict(zip(headers, values)))
        tables.append({"headers": headers, "rows": rows, "so_dong": len(rows)})

    return paragraphs, tables


def tao_output(paragraphs, tables, duong_dan: Path):
    return {
        "trang_thai": "success",
        "ten_file": duong_dan.name,
        "loai_file": "docx",
        "pdf_mode": None,
        "ocr_da_su_dung": False,
        "du_lieu": {
            "paragraphs": paragraphs,
            "so_bang": len(tables),
            "bang": tables,
        },
        "loi": None,
    }


def main():
    if not TEST_FILE.exists():
        raise FileNotFoundError(f"Không tìm thấy {TEST_FILE}. Hãy chạy generate_test_data.py trước.")

    paragraphs, tables = doc_docx(TEST_FILE)
    print(f"File       : {TEST_FILE}")
    print(f"Paragraphs : {len(paragraphs)}")
    print(f"Số bảng    : {len(tables)}")
    if tables:
        print(f"Headers    : {tables[0]['headers']}")
        print(f"Số record  : {tables[0]['so_dong']}")

    print("\nFILE READER OUTPUT")
    print(json.dumps(tao_output(paragraphs, tables, TEST_FILE), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
