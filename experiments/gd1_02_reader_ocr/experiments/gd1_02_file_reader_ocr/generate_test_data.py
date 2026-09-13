# -*- coding: utf-8 -*-
"""
GĐ1-02: Tạo dữ liệu test thống nhất cho File Reader + OCR.

Mục tiêu:
- Tạo cùng một bộ dữ liệu nhân viên ở 5 định dạng: XLSX, CSV, DOCX,
  PDF text và PDF scan.
- Dữ liệu chỉ phục vụ kiểm thử File Reader/OCR, không phải dữ liệu train AI.
- File Reader phải đọc động theo header, không hard-code số lượng cột.
- Tất cả định dạng phải dùng đúng cùng một bộ record và cùng schema test.
"""

from __future__ import annotations

import csv
import os
import random
from pathlib import Path

THU_MUC_HIEN_TAI = Path(__file__).resolve().parent
THU_MUC_DU_AN = THU_MUC_HIEN_TAI.parents[1]
THU_MUC_TEST = THU_MUC_DU_AN / "test_data" / "gd1_02"
THU_MUC_TEST.mkdir(parents=True, exist_ok=True)

# Schema test chuẩn của dữ liệu nhân viên trong GĐ1-02.
# Không bắt buộc mã đơn vị vì hệ thống hiện không dùng bảng đơn vị riêng.
DU_LIEU_MAU = [
    {
        "ma_nhan_vien": "NV001",
        "ho_ten": "Nguyễn Văn An",
        "don_vi": "Trung tâm Công nghệ thông tin",
        "chuc_vu": "Chuyên viên",
        "email": "nv001@example.com",
    },
    {
        "ma_nhan_vien": "NV002",
        "ho_ten": "Trần Thị Bích Ngọc",
        "don_vi": "Phòng Kế hoạch Tài chính",
        "chuc_vu": "Kế toán viên",
        "email": "nv002@example.com",
    },
    {
        "ma_nhan_vien": "NV003",
        "ho_ten": "Lê Hoàng Phúc",
        "don_vi": "Trung tâm Công nghệ thông tin",
        "chuc_vu": "Kỹ sư phần mềm",
        "email": "nv003@example.com",
    },
    {
        "ma_nhan_vien": "NV004",
        "ho_ten": "Phạm Minh Đức",
        "don_vi": "Ban Quản lý Dự án",
        "chuc_vu": "Chuyên viên dự án",
        "email": "nv004@example.com",
    },
    {
        "ma_nhan_vien": "NV005",
        "ho_ten": "Võ Thị Hồng Nhung",
        "don_vi": "Chi nhánh Đà Nẵng",
        "chuc_vu": "Nhân viên hành chính",
        "email": "nv005@example.com",
    },
    {
        "ma_nhan_vien": "NV006",
        "ho_ten": "Đỗ Quốc Khánh",
        "don_vi": "Phòng Tổ chức Nhân sự",
        "chuc_vu": "Chuyên viên nhân sự",
        "email": "nv006@example.com",
    },
    {
        "ma_nhan_vien": "NV007",
        "ho_ten": "Huỳnh Thanh Tâm",
        "don_vi": "Phòng Kế hoạch Tài chính",
        "chuc_vu": "Chuyên viên tài chính",
        "email": "nv007@example.com",
    },
    {
        "ma_nhan_vien": "NV008",
        "ho_ten": "Ngô Bảo Trâm",
        "don_vi": "Văn phòng Đại diện Hà Nội",
        "chuc_vu": "Chuyên viên",
        "email": "nv008@example.com",
    },
]

TIEU_DE = ["Mã NV", "Họ và tên", "Đơn vị", "Chức vụ", "Email"]
KHOA_DU_LIEU = ["ma_nhan_vien", "ho_ten", "don_vi", "chuc_vu", "email"]


def tim_font_tieng_viet() -> str | None:
    """Tìm font Unicode phổ biến trên Windows/Linux/macOS."""
    ung_vien = [
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/times.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for duong_dan in ung_vien:
        if os.path.exists(duong_dan):
            return duong_dan
    return None


def tao_xlsx() -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    duong_dan = THU_MUC_TEST / "test_nhan_vien.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách nhân viên"

    for cot, tieu_de in enumerate(TIEU_DE, 1):
        o = ws.cell(row=1, column=cot, value=tieu_de)
        o.font = Font(bold=True)
        o.alignment = Alignment(horizontal="center")
        o.border = Border(bottom=Side(style="thin"))

    for dong, ban_ghi in enumerate(DU_LIEU_MAU, 2):
        for cot, khoa in enumerate(KHOA_DU_LIEU, 1):
            ws.cell(row=dong, column=cot, value=ban_ghi[khoa])

    for cot in ws.columns:
        do_rong = max(len(str(o.value or "")) for o in cot) + 3
        ws.column_dimensions[cot[0].column_letter].width = min(do_rong, 40)

    wb.save(duong_dan)
    return duong_dan


def tao_csv() -> Path:
    duong_dan = THU_MUC_TEST / "test_nhan_vien.csv"
    with open(duong_dan, "w", encoding="utf-8-sig", newline="") as tep:
        writer = csv.writer(tep)
        writer.writerow(TIEU_DE)
        for ban_ghi in DU_LIEU_MAU:
            writer.writerow([ban_ghi[k] for k in KHOA_DU_LIEU])
    return duong_dan


def tao_docx() -> Path:
    from docx import Document

    duong_dan = THU_MUC_TEST / "test_nhan_vien.docx"
    doc = Document()
    doc.add_heading("DANH SÁCH NHÂN VIÊN", level=1)
    doc.add_paragraph("Dữ liệu giả lập dùng để kiểm thử File Reader – GĐ1-02.")

    bang = doc.add_table(rows=1, cols=len(TIEU_DE))
    bang.style = "Table Grid"
    for i, tieu_de in enumerate(TIEU_DE):
        bang.rows[0].cells[i].text = tieu_de

    for ban_ghi in DU_LIEU_MAU:
        cells = bang.add_row().cells
        for i, khoa in enumerate(KHOA_DU_LIEU):
            cells[i].text = str(ban_ghi[khoa])

    doc.add_paragraph("Ghi chú: Toàn bộ thông tin trong file là dữ liệu giả lập.")
    doc.save(duong_dan)
    return duong_dan


def tao_pdf_text() -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    duong_dan = THU_MUC_TEST / "test_nhan_vien_text.pdf"
    font_name = "Helvetica"
    font_path = tim_font_tieng_viet()
    if font_path:
        try:
            pdfmetrics.registerFont(TTFont("FontVN", font_path))
            font_name = "FontVN"
        except Exception:
            pass

    doc = SimpleDocTemplate(
        str(duong_dan),
        pagesize=landscape(A4),
        leftMargin=1.0 * cm,
        rightMargin=1.0 * cm,
        topMargin=1.0 * cm,
        bottomMargin=1.0 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TieuDeVN", parent=styles["Title"], fontName=font_name, fontSize=16
    )
    note_style = ParagraphStyle(
        "NoiDungVN", parent=styles["Normal"], fontName=font_name, fontSize=9
    )

    du_lieu_bang = [TIEU_DE] + [
        [ban_ghi[k] for k in KHOA_DU_LIEU] for ban_ghi in DU_LIEU_MAU
    ]
    bang = Table(du_lieu_bang, repeatRows=1)
    bang.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )

    doc.build(
        [
            Paragraph("DANH SÁCH NHÂN VIÊN", title_style),
            Spacer(1, 0.3 * cm),
            Paragraph("PDF text-based dùng để kiểm thử GĐ1-02.", note_style),
            Spacer(1, 0.4 * cm),
            bang,
        ]
    )
    return duong_dan


def tao_pdf_scan() -> Path:
    """Tạo PDF chỉ chứa ảnh để mô phỏng tài liệu scan."""
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    duong_dan = THU_MUC_TEST / "test_nhan_vien_scan.pdf"
    anh_tam = THU_MUC_TEST / "_temp_scan_page.png"

    dpi = 180
    width_px = int(11.69 * dpi)
    height_px = int(8.27 * dpi)
    img = Image.new("RGB", (width_px, height_px), "white")
    draw = ImageDraw.Draw(img)

    font_path = tim_font_tieng_viet()
    if font_path:
        font_title = ImageFont.truetype(font_path, 28)
        font_header = ImageFont.truetype(font_path, 16)
        font_body = ImageFont.truetype(font_path, 14)
    else:
        font_title = font_header = font_body = ImageFont.load_default()

    draw.text((60, 45), "DANH SÁCH NHÂN VIÊN", fill="black", font=font_title)
    draw.text((60, 90), "PDF scan giả lập – GĐ1-02", fill="gray", font=font_body)

    col_widths = [115, 280, 390, 260, 300]
    x_start = 50
    y_start = 135
    row_height = 38

    x = x_start
    for i, tieu_de in enumerate(TIEU_DE):
        draw.rectangle(
            [x, y_start, x + col_widths[i], y_start + row_height],
            outline="black",
            fill="#4472C4",
        )
        draw.text((x + 4, y_start + 9), tieu_de, fill="white", font=font_header)
        x += col_widths[i]

    for row_idx, ban_ghi in enumerate(DU_LIEU_MAU):
        y = y_start + (row_idx + 1) * row_height
        x = x_start
        bg = "#FFFFFF" if row_idx % 2 == 0 else "#EAF0F8"
        for col_idx, khoa in enumerate(KHOA_DU_LIEU):
            draw.rectangle(
                [x, y, x + col_widths[col_idx], y + row_height],
                outline="black",
                fill=bg,
            )
            draw.text(
                (x + 4, y + 9),
                str(ban_ghi[khoa]),
                fill="black",
                font=font_body,
            )
            x += col_widths[col_idx]

    # Nhiễu nhẹ, cố định seed để lần chạy nào cũng tạo cùng kiểu ảnh test.
    random.seed(42)
    for _ in range(250):
        rx = random.randint(0, width_px - 1)
        ry = random.randint(0, height_px - 1)
        draw.point((rx, ry), fill=(205, 205, 205))

    img.save(anh_tam, "PNG")

    page_w, page_h = landscape(A4)
    c = canvas.Canvas(str(duong_dan), pagesize=(page_w, page_h))
    c.drawImage(str(anh_tam), 0, 0, width=page_w, height=page_h)
    c.showPage()
    c.save()

    try:
        anh_tam.unlink()
    except OSError:
        pass

    return duong_dan


def main() -> None:
    cac_file = [
        tao_xlsx(),
        tao_csv(),
        tao_docx(),
        tao_pdf_text(),
        tao_pdf_scan(),
    ]

    print("=" * 72)
    print("ĐÃ TẠO BỘ DỮ LIỆU TEST GĐ1-02")
    print("=" * 72)
    for tep in cac_file:
        print(f"[OK] {tep.name} ({tep.stat().st_size} bytes)")
    print(f"Thư mục: {THU_MUC_TEST}")


if __name__ == "__main__":
    main()
