# -*- coding: utf-8 -*-
"""GĐ1-02: Thử nghiệm đọc XLSX bằng openpyxl và pandas."""

from __future__ import annotations

import json
from pathlib import Path

THU_MUC_DU_AN = Path(__file__).resolve().parents[2]
TEST_FILE = THU_MUC_DU_AN / "test_data" / "gd1_02" / "test_nhan_vien.xlsx"


def doc_bang_openpyxl(duong_dan: Path):
    from openpyxl import load_workbook

    wb = load_workbook(duong_dan, read_only=True, data_only=True)
    ws = wb.active
    cac_dong = list(ws.iter_rows(values_only=True))
    wb.close()

    if not cac_dong:
        return [], []

    headers = [str(v).strip() if v is not None else "" for v in cac_dong[0]]
    rows = []
    for dong in cac_dong[1:]:
        if not any(v not in (None, "") for v in dong):
            continue
        rows.append(dict(zip(headers, dong)))
    return headers, rows


def doc_bang_pandas(duong_dan: Path):
    try:
        import pandas as pd
    except ImportError:
        return None

    df = pd.read_excel(duong_dan, engine="openpyxl")
    return df.to_dict(orient="records")


def tao_output(headers, rows, duong_dan: Path):
    return {
        "trang_thai": "success",
        "ten_file": duong_dan.name,
        "loai_file": "xlsx",
        "pdf_mode": None,
        "ocr_da_su_dung": False,
        "du_lieu": {
            "headers": headers,
            "rows": rows,
            "so_dong": len(rows),
        },
        "loi": None,
    }


def main():
    if not TEST_FILE.exists():
        raise FileNotFoundError(f"Không tìm thấy {TEST_FILE}. Hãy chạy generate_test_data.py trước.")

    headers, rows = doc_bang_openpyxl(TEST_FILE)
    records_pandas = doc_bang_pandas(TEST_FILE)

    print(f"File: {TEST_FILE}")
    print(f"Headers: {headers}")
    print(f"Số record openpyxl: {len(rows)}")
    if records_pandas is not None:
        print(f"Số record pandas  : {len(records_pandas)}")
        print(f"Khớp số dòng      : {len(rows) == len(records_pandas)}")

    print("\nFILE READER OUTPUT")
    print(json.dumps(tao_output(headers, rows, TEST_FILE), ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
